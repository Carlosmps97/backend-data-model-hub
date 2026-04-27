"""Herramientas del KnowledgeBaseAgent — carga y consulta de lineamientos de modelamiento.

Soporta múltiples formatos de fuente: JSON, XLSX, DOCX, PDF, MD, TXT.

Dos fuentes de guidelines coexisten:
- **Por sesión** (`_guidelines_cache_by_session`): contenido cargado desde la
  API para una `conversation_id` específica. Tiene prioridad si el caller
  marca la sesión activa con `session_scope(session_id)` o
  `set_active_session(session_id)`.
- **Global** (`_guidelines_cache_global`): contenido del archivo configurado
  vía `settings.GUIDELINES_PATH`. Se usa como fallback cuando no hay sesión
  activa (CLI, tests, ejecuciones directas).

Las firmas `@tool` de `query_guidelines` y `get_all_guidelines` se mantienen
idénticas: el LLM no debe conocer `session_id`. La sesión activa se propaga
internamente vía `ContextVar`.
"""

from __future__ import annotations

import json
from contextlib import contextmanager
from contextvars import ContextVar
from pathlib import Path
from typing import Annotated, Iterator

from agent_framework import tool
from pydantic import Field

from src.config import settings

# ─── Estado interno ─────────────────────────────────────────────────────

# Cache global (para compatibilidad con CLI y ejecuciones directas).
_guidelines_cache_global: dict | None = None

# Cache por sesión: conversation_id -> guidelines procesadas.
_guidelines_cache_by_session: dict[str, dict] = {}

# ContextVar que indica qué sesión está activa en el contexto async actual.
# `None` => usar cache/path global.
_active_session: ContextVar[str | None] = ContextVar(
    "data_modeler_active_session", default=None
)


# ─── Helpers de sesión activa ──────────────────────────────────────────


def set_active_session(session_id: str | None) -> object:
    """Marca la sesión activa para el contexto async actual.

    Retorna el token devuelto por `ContextVar.set` para que el caller pueda
    restaurar el estado anterior con `reset_active_session(token)`.
    """
    return _active_session.set(session_id)


def reset_active_session(token: object) -> None:
    """Restaura el estado de la sesión activa al valor previo al token."""
    _active_session.reset(token)  # type: ignore[arg-type]


@contextmanager
def session_scope(session_id: str | None) -> Iterator[None]:
    """Context manager para fijar la sesión activa durante un bloque."""
    token = set_active_session(session_id)
    try:
        yield
    finally:
        reset_active_session(token)


# ─── Carga y procesamiento de archivos de guidelines ────────────────────


def _process_guidelines_path(path: Path) -> dict:
    """Procesa un archivo de guidelines y devuelve la estructura cargada.

    Soporta JSON, XLSX, DOCX, PDF, MD, TXT. Retorna un dict ya consumible
    por `query_guidelines` y `get_all_guidelines`.
    """
    if not path.exists():
        return {"error": f"No se encontró el archivo de lineamientos: {path}"}

    suffix = path.suffix.lower()

    if suffix == ".json":
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)

    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        data: dict[str, list[dict]] = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[dict] = []
            headers: list[str] | None = None
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [
                        str(c).strip().lower() if c else f"col_{j}"
                        for j, c in enumerate(row)
                    ]
                else:
                    rows.append(dict(zip(headers or [], row)))
            data[sheet_name] = rows
        wb.close()
        return {"format": "xlsx", "sheets": data}

    if suffix in (".docx", ".pdf"):
        # Convertir a Markdown con Docling para preservar estructura.
        from src.tools.convert_tools import _convert_document

        try:
            md_path = _convert_document(path)
            text = md_path.read_text(encoding="utf-8")
            return {"format": "markdown", "content": text}
        except Exception as e:
            return {"error": f"Error al convertir {suffix} a Markdown: {e}"}

    if suffix in (".md", ".txt", ".text"):
        text = path.read_text(encoding="utf-8")
        fmt = "markdown" if suffix == ".md" else "text"
        return {"format": fmt, "content": text}

    return {"error": f"Formato no soportado: {suffix}"}


def process_guidelines_file(path: Path | str) -> dict:
    """API pública: procesa un archivo de guidelines a un dict consumible.

    Útil para el endpoint de upload de la API: recibe un archivo subido por
    el usuario, lo guarda en disco, y este helper lo procesa.
    """
    return _process_guidelines_path(Path(path))


# ─── Cache de sesión ────────────────────────────────────────────────────


def set_session_guidelines(session_id: str, guidelines: dict) -> None:
    """Asocia un contenido de guidelines ya procesado a una sesión."""
    _guidelines_cache_by_session[session_id] = guidelines


def get_session_guidelines(session_id: str) -> dict | None:
    """Recupera el contenido de guidelines asociado a una sesión, si existe."""
    return _guidelines_cache_by_session.get(session_id)


def clear_session_guidelines(session_id: str) -> None:
    """Elimina las guidelines de una sesión del cache (idempotente)."""
    _guidelines_cache_by_session.pop(session_id, None)


# ─── Resolución del cache aplicable al contexto actual ─────────────────


def _load_guidelines_global() -> dict:
    """Carga (y cachea) las guidelines desde la ruta global configurada."""
    global _guidelines_cache_global
    if _guidelines_cache_global is not None:
        return _guidelines_cache_global
    _guidelines_cache_global = _process_guidelines_path(
        Path(settings.GUIDELINES_PATH)
    )
    return _guidelines_cache_global


def _resolve_guidelines() -> dict:
    """Resuelve qué guidelines aplican al contexto actual.

    Si hay sesión activa y tiene guidelines cargadas, las retorna.
    Si hay sesión activa pero no tiene guidelines, retorna las globales como
    fallback (esto permite a los agentes seguir funcionando aunque la sesión
    no haya cargado un archivo propio).
    Sin sesión activa, retorna las globales.
    """
    session_id = _active_session.get()
    if session_id is not None:
        session_data = _guidelines_cache_by_session.get(session_id)
        if session_data is not None:
            return session_data
    return _load_guidelines_global()


# ─── Tools expuestas a los agentes ──────────────────────────────────────


@tool(approval_mode="never_require")
def query_guidelines(
    topic: Annotated[
        str,
        Field(
            description=(
                "Tema a consultar en los lineamientos de modelamiento. "
                "Ejemplos: 'naming_conventions', 'audit_columns', 'data_type_mappings', "
                "'primary_key_conventions', 'foreign_key_conventions', 'general_rules'. "
                "También puede ser un motor de BD como 'databricks_sql', 'sqlserver', etc."
            )
        ),
    ],
) -> str:
    """Consulta los lineamientos corporativos de modelamiento de datos por tema específico."""
    guidelines = _resolve_guidelines()

    if "error" in guidelines:
        return json.dumps(guidelines)

    # Si es JSON estructurado, buscar por clave
    if isinstance(guidelines, dict) and "format" not in guidelines:
        topic_lower = topic.lower().strip()

        # Buscar directamente en las claves del JSON
        if topic_lower in guidelines:
            return json.dumps(guidelines[topic_lower], indent=2, ensure_ascii=False)

        # Buscar en data_type_mappings por motor
        if topic_lower in guidelines.get("data_type_mappings", {}):
            return json.dumps(
                {"engine": topic_lower, "type_mappings": guidelines["data_type_mappings"][topic_lower]},
                indent=2,
                ensure_ascii=False,
            )

        # Búsqueda por keywords en todas las claves
        results = {}
        for key, value in guidelines.items():
            if topic_lower in key.lower() or (
                isinstance(value, (dict, list))
                and topic_lower in json.dumps(value, ensure_ascii=False).lower()
            ):
                results[key] = value

        if results:
            return json.dumps(results, indent=2, ensure_ascii=False)

        return json.dumps(
            {"message": f"No se encontraron lineamientos para el tema: {topic}"},
            ensure_ascii=False,
        )

    # Si es texto plano (de docx, pdf, md, txt)
    if "content" in guidelines:
        content = guidelines["content"]
        lines = content.split("\n")
        relevant = [l for l in lines if topic.lower() in l.lower()]
        if relevant:
            return "\n".join(relevant)
        return f"No se encontró información específica sobre '{topic}' en los lineamientos."

    return json.dumps(guidelines, indent=2, ensure_ascii=False)


@tool(approval_mode="never_require")
def get_all_guidelines() -> str:
    """Retorna todos los lineamientos corporativos de modelamiento de datos completos."""
    guidelines = _resolve_guidelines()
    return json.dumps(guidelines, indent=2, ensure_ascii=False)
