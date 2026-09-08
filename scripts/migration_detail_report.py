"""Reporte Excel de INCONGRUENCIAS de la migración Erwin, objeto por objeto.

Enfocado en lo que los modeladores deben corregir EN ERWIN: la hoja
«Incongruencias» resume cada tipo de hallazgo en los archivos XML; las demás
hojas listan caso por caso, con lo que la migración hizo al respecto y la
ubicación exacta en la plataforma (proyecto / carpeta / canvas) para poder
revisar cada caso. Las uniones limpias (la misma tabla con las mismas columnas
en varios archivos) NO se reportan: no son incongruencia.

Fuentes (las tres se cruzan):
  1. `migration-reports/` — el último `run_migration-*.json` y el último
     reporte por XML de esa corrida (decisiones tabla por tabla).
  2. La base Lakebase — estado FINAL (post `audit_data_consistency --fix`).
  3. Los XML de Erwin — estructuras por archivo (columnas, particiones,
     definiciones UDP, diagramas) para detectar lo que no conversa entre sí.

Uso (desde backend-data-model-hub/):
    .venv/bin/python -m scripts.migration_detail_report
    .venv/bin/python -m scripts.migration_detail_report \
        --xml-root ../folder_data --out ../REPORTE-MIGRACION-ERWIN-DETALLE.xlsx

`--cache-dir` guarda los modelos parseados (pickle) para re-ejecutar rápido.
Requiere `openpyxl` (está en el venv local; NO es dependencia de la app).
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import pickle
import sys
from collections import defaultdict

from dotenv import load_dotenv

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)
load_dotenv(os.path.join(BACKEND, ".env"))

from app.core.db.sync import get_sync_db  # noqa: E402
from scripts.erwin_migration import erwin_parser as ep  # noqa: E402
from scripts.erwin_migration import policies as pol  # noqa: E402

ACTIVE = {"flgactive": {"$ne": False}}
MAX_PATHS = 8
MAX_NAMES = 12
REL_TYPE_ES = {ep.REL_IDENTIFYING: "identificante", ep.REL_NON_IDENTIFYING: "no identificante",
               ep.REL_SUBTYPE: "subcategoría (supertipo→subtipo)", ep.REL_TABLE_TO_VIEW: "tabla→vista"}
NIVEL_ES = {"table": "Tabla", "column": "Columna", "canvas": "Canvas"}


def log(msg: str) -> None:
    print(msg, flush=True)


def cap_names(names, limit=MAX_NAMES):
    names = sorted(names)
    if len(names) <= limit:
        return ", ".join(names)
    return ", ".join(names[:limit]) + f" … (+{len(names) - limit} más)"


# ═══════════════════════════ descubrimiento de fuentes ═════════════════════
def discover(reports_dir: str, xml_root: str) -> list[dict]:
    """[{tag, project, domain, xml_path, report}] desde el último run_migration."""
    runs = sorted(glob.glob(os.path.join(reports_dir, "run_migration-*.json")))
    if not runs:
        raise SystemExit(f"No hay run_migration-*.json en {reports_dir}")
    with open(runs[-1]) as f:
        run = json.load(f)
    out = []
    for entry in run["files"]:
        rel = entry["rel"]
        base = os.path.basename(rel)
        pat = os.path.join(reports_dir, base.replace(" ", "_") + "-*.json")
        reps = sorted(glob.glob(pat))
        if not reps:
            raise SystemExit(f"No hay reporte por archivo para {base} ({pat})")
        with open(reps[-1]) as f:
            report = json.load(f)
        project, domain = entry["project"], entry["domain"]
        tag = f"{project.replace('Modelo ', '')} {domain.replace('Modelo ', '')}"
        out.append({
            "tag": tag, "project": project, "domain": domain, "report": report,
            "xml_path": os.path.join(xml_root, rel),
            "run": run,
        })
    return out


def parse_models(files: list[dict], cache_dir: str | None) -> None:
    for fi in files:
        path = fi["xml_path"]
        cache = None
        if cache_dir:
            os.makedirs(cache_dir, exist_ok=True)
            st = os.stat(path)
            cache = os.path.join(
                cache_dir, f"{os.path.basename(path)}.{st.st_size}.{int(st.st_mtime)}.pkl")
        if cache and os.path.exists(cache):
            with open(cache, "rb") as f:
                fi["model"] = pickle.load(f)
            log(f"[xml] {fi['tag']}: cache")
            continue
        log(f"[xml] {fi['tag']}: parseando {os.path.getsize(path) // 1024 // 1024} MB …")
        fi["model"] = ep.parse(path)
        if cache:
            with open(cache, "wb") as f:
                pickle.dump(fi["model"], f)


# ═══════════════════════════ extracción ════════════════════════════════════
def extract(files: list[dict]) -> dict:
    models = {fi["tag"]: fi["model"] for fi in files}
    reports = {fi["tag"]: fi["report"] for fi in files}
    proj_of_tag = {fi["tag"]: fi["project"] for fi in files}
    dom_of_tag = {fi["tag"]: fi["domain"] for fi in files}

    schema_of = {tag: m.owner_schema() for tag, m in models.items()}
    ent_file: dict[str, list[str]] = {}
    rel_file: dict[str, str] = {}
    view_file: dict[str, list[str]] = {}
    for tag, m in models.items():
        for e in m.entities.values():
            ent_file.setdefault(e.id, []).append(tag)
        for v in m.views.values():
            view_file.setdefault(v.id, []).append(tag)
        for r in m.relationships.values():
            rel_file[r.id] = tag

    rel_diags = {tag: defaultdict(list) for tag in models}
    obj_diags = {tag: defaultdict(list) for tag in models}
    for tag, m in models.items():
        rel_ids, known = set(m.relationships), set(m.entities) | set(m.views)
        for d in m.diagrams:
            for ref, _a in d.shapes:
                if ref in rel_ids:
                    rel_diags[tag][ref].append((d.subject_area, d.name))
                elif ref in known:
                    obj_diags[tag][ref].append((d.subject_area, d.name))

    log("[xml] definiciones y valores UDP por archivo …")
    collapsed_by_tag = {tag: pol.udp_defs_by_view(m.udp_defs) for tag, m in models.items()}
    udp_vals_by_tag = {tag: pol.resolve_udp_values(m.udp_values, collapsed_by_tag[tag])
                       for tag, m in models.items()}

    db = get_sync_db()
    log("[db] proyectos/carpetas/canvases …")
    projects = {p["_id"]: p.get("name") for p in db.projects.find(ACTIVE, {"name": 1})}
    folders = {f["_id"]: f for f in db.folders.find(
        ACTIVE, {"name": 1, "parentFolderId": 1, "projectId": 1})}
    sas = {s["_id"]: s for s in db.subject_areas.find(
        {}, {"name": 1, "folderId": 1, "projectId": 1, "tableIds": 1, "layout": 1,
             "erwinLongId": 1, "udpValues": 1})}

    log("[db] tablas / columnas / relaciones / vistas …")
    tables = {t["_id"]: t for t in db.canonical_tables.find(
        {}, {"schema": 1, "physicalName": 1, "logicalName": 1, "erwinLongId": 1,
             "flgactive": 1, "udpValues": 1})}
    t_active = {tid for tid, t in tables.items() if t.get("flgactive") is not False}
    by_phys, by_key, by_erwin_t = {}, {}, {}
    for tid in t_active:
        t = tables[tid]
        by_phys[(t.get("physicalName") or "").upper()] = tid
        by_key[((t.get("schema") or "").upper(), (t.get("physicalName") or "").upper())] = tid
        if t.get("erwinLongId"):
            by_erwin_t[t["erwinLongId"]] = tid

    cols = {c["_id"]: c for c in db.canonical_columns.find(
        {}, {"tableId": 1, "physicalName": 1, "logicalName": 1, "isPrimaryKey": 1,
             "pkPosition": 1, "ordinal": 1, "flgactive": 1, "deletedAt": 1, "erwinLongId": 1})}
    c_active = {cid for cid, c in cols.items() if c.get("flgactive") is not False}
    act_colnames: dict[str, set] = defaultdict(set)
    pk_of_table: dict[str, list] = defaultdict(list)
    for cid in c_active:
        c = cols[cid]
        act_colnames[c["tableId"]].add((c.get("physicalName") or "").upper())
        if c.get("isPrimaryKey"):
            pk_of_table[c["tableId"]].append(
                (c.get("pkPosition") if c.get("pkPosition") is not None else 999, cid))
    for tid in pk_of_table:
        pk_of_table[tid].sort()

    rels = {r["_id"]: r for r in db.relationships.find({}, {
        "parentTableId": 1, "childTableId": 1, "pairs": 1, "identifying": 1,
        "subcategory": 1, "flgactive": 1, "deletedAt": 1, "erwinLongId": 1})}
    r_active = {rid for rid, r in rels.items() if r.get("flgactive") is not False}
    views = {v["_id"]: v for v in db.views.find(ACTIVE, {
        "name": 1, "schema": 1, "sourceTableIds": 1, "sources": 1, "joinOverride": 1,
        "erwinLongId": 1})}
    udp_defs = {d["_id"]: d for d in db.udp_definitions.find(ACTIVE)}

    # ── ubicación ──────────────────────────────────────────────────────────
    def folder_chain(fid):
        names, seen = [], set()
        while fid and fid in folders and fid not in seen:
            seen.add(fid)
            names.append(folders[fid].get("name") or "?")
            fid = folders[fid].get("parentFolderId")
        return list(reversed(names))

    sa_path = {}
    for sid, s in sas.items():
        sa_path[sid] = (projects.get(s.get("projectId"), "?"),
                        " / ".join(folder_chain(s.get("folderId")) + [s.get("name") or "?"]))

    canvases_of_table: dict[str, list] = defaultdict(list)
    layout_of_node: dict[str, list] = defaultdict(list)
    for sid, s in sas.items():
        for tid in (s.get("tableIds") or []):
            canvases_of_table[tid].append(sid)
        for nid in (s.get("layout") or {}):
            layout_of_node[nid].append(sid)

    def paths_of(said_list):
        return sorted({sa_path[sid][1] for sid in said_list if sid in sa_path})

    def cap_paths(paths, empty="(no aparece en ningún canvas)"):
        if not paths:
            return empty
        if len(paths) <= MAX_PATHS:
            return "\n".join(paths)
        return "\n".join(paths[:MAX_PATHS] +
                         [f"… (+{len(paths) - MAX_PATHS} canvases más — total {len(paths)})"])

    def project_of(said_list, fallback=""):
        ps = sorted({sa_path[sid][0] for sid in said_list if sid in sa_path})
        return " y ".join(ps) if ps else fallback

    def table_loc(tid, fallback_proj=""):
        cs = canvases_of_table.get(tid, [])
        return project_of(cs, fallback_proj), cap_paths(paths_of(cs))

    def rel_canvases(parent_tid, child_tid):
        both = set(canvases_of_table.get(parent_tid, [])) & set(canvases_of_table.get(child_tid, []))
        return project_of(both), paths_of(both)

    def colname(cid, with_dead_mark=True):
        c = cols.get(cid)
        if not c:
            return f"(columna inexistente {str(cid)[:8]})"
        n = c.get("physicalName") or "?"
        if with_dead_mark and c.get("flgactive") is False:
            n += " (no vigente)"
        return n

    def tname(tid):
        t = tables.get(tid)
        return f"{t.get('schema')}.{t.get('physicalName')}" if t else f"(tabla inexistente {str(tid)[:8]})"

    def tlogical(tid):
        t = tables.get(tid)
        return (t.get("logicalName") or "") if t else ""

    OUT: dict = {}

    # ═══ Tablas homónimas (_DUPn) ══════════════════════════════════════════
    rows = []
    for tag in models:
        proj = proj_of_tag[tag]
        for g in reports[tag]["decisions"]["renamed_dups"]:
            schema_raw, _orig = g["key"].split(".", 1)
            kept_tid = by_phys.get((g["kept"] or "").upper())
            _kp, kept_canv = table_loc(kept_tid, proj) if kept_tid else (proj, "?")
            for r in g["renamed"]:
                tid = by_phys.get(r["to"].upper())
                t = tables.get(tid, {})
                p, canv = table_loc(tid, proj) if tid else (proj, "?")
                rows.append({
                    "archivo": tag, "proyecto": p or proj,
                    "esquema": t.get("schema") or schema_raw,
                    "nombre_original": r["from"], "nombre_final": r["to"],
                    "nombre_logico": r.get("logicalName") or t.get("logicalName") or "",
                    "copias": g["copies"], "conservo_nombre": g["kept"],
                    "canvases": canv, "canvases_kept": kept_canv,
                })
    OUT["homonimas"] = rows
    log(f"[homónimas] copias renombradas _DUPn: {len(rows)}")

    # ═══ Columnas distintas entre archivos (misma tabla, merge pendiente) ══
    # Contribuyentes por tabla final: (1) por GUID Erwin, (2) por clave natural
    # de las adopciones y conflictos registrados en los reportes de la corrida,
    # (3) por el GUID de cada columna no vigente (ancla al archivo cuya versión
    # de la tabla no quedó, aunque el GUID de tabla haya cambiado al ganar el
    # otro archivo).
    contrib: dict[str, dict[str, dict]] = defaultdict(dict)   # tid → tag → {COL: (attr, es_pk)}

    def add_contrib(tid, tag, e):
        slot = contrib[tid].setdefault(tag, {})
        for a in e.attributes:
            if a.physical:
                slot.setdefault(a.physical.upper(), (a, a.id in e.pk_attr_ids))

    for tag, m in models.items():
        for e in m.entities.values():
            tid = by_erwin_t.get(e.id)
            if tid:
                add_contrib(tid, tag, e)
    for tag, m in models.items():
        sof = schema_of[tag]
        ent_by_key = {f"{pol.schema_or_default(sof.get(e.id))}.{e.physical}".upper(): e
                      for e in m.entities.values() if e.physical}
        dec = reports[tag]["decisions"]
        for a in dec["adopted"] + dec["conflicts"]:
            e = ent_by_key.get(a["key"].upper())
            if not e:
                continue
            s, p = a["key"].upper().split(".", 1)
            tid = by_key.get((s, p)) or by_phys.get(p)
            if tid:
                add_contrib(tid, tag, e)
    attr_owner = {}
    for tag, m in models.items():
        attr_owner[tag] = {aid: a.owner_id for aid, a in m.attr_index().items()}
    for cid, c in cols.items():
        if c.get("flgactive") is not False:
            continue
        aid, tid = c.get("erwinLongId"), c.get("tableId")
        if not aid or not tid or tid not in t_active:
            continue
        for tag, m in models.items():
            oid = attr_owner[tag].get(aid)
            if oid and oid in m.entities:
                add_contrib(tid, tag, m.entities[oid])

    rows, n_no_vig, n_tablas_diff = [], 0, 0
    for tid, per_tag in contrib.items():
        if len(per_tag) < 2:
            continue
        sets = {tag: set(d) for tag, d in per_tag.items()}
        union = set().union(*sets.values())
        inter = set.intersection(*sets.values())
        if union == inter:
            continue                       # unión limpia: no se reporta
        n_tablas_diff += 1
        t = tables[tid]
        p, canv = table_loc(tid)
        for cn in sorted(union - inter):
            present = sorted(tag for tag in sets if cn in sets[tag])
            absent = sorted(tag for tag in sets if cn not in sets[tag])
            a, es_pk = per_tag[present[0]][cn]
            vigente = cn in act_colnames.get(tid, set())
            if not vigente:
                n_no_vig += 1
            rows.append({
                "proyecto": p, "esquema": t.get("schema") or "",
                "tabla": t.get("physicalName") or "",
                "columna": a.physical, "columna_logica": a.name or "",
                "es_pk": "Sí" if es_pk else "No",
                "presente": ", ".join(present), "falta": ", ".join(absent),
                "estado": ("Vigente (la versión que quedó la incluye)" if vigente else
                           "No vigente (la versión que quedó NO la incluye)"),
                "canvases": canv,
            })
    rows.sort(key=lambda x: (x["tabla"], x["columna"]))
    OUT["cols_distintas"] = rows
    n_dead = sum(1 for c in cols.values() if c.get("flgactive") is False)
    log(f"[cols distintas] {len(rows)} columnas en {n_tablas_diff} tablas "
        f"(no vigentes: {n_no_vig}; columnas no vigentes en BD: {n_dead})")

    # ═══ Columnas duplicadas (las únicas realmente eliminadas) ═════════════
    rows = []
    for tag in models:
        proj = proj_of_tag[tag]
        for d in reports[tag]["decisions"]["cols_dropped"]:
            schema_raw, phys = d["table"].split(".", 1)
            tid = by_key.get((schema_raw.upper(), phys.upper())) or by_phys.get(phys.upper())
            t = tables.get(tid, {})
            p, canv = table_loc(tid, proj) if tid else (proj, "?")
            rows.append({"archivo": tag, "proyecto": p or proj,
                         "esquema": t.get("schema") or schema_raw, "tabla": phys,
                         "columna": d["column"], "canvases": canv})
    OUT["cols_duplicadas"] = rows
    log(f"[cols duplicadas] {len(rows)}")

    # ═══ Relaciones incongruentes (una sola hoja) ══════════════════════════
    rows = []
    #  a) vínculo con columna no vigente (desactivadas por la corrección)
    for rid, r in rels.items():
        if r.get("flgactive") is not False:
            continue
        pt, ct = r.get("parentTableId"), r.get("childTableId")
        proj, canvs = rel_canvases(pt, ct)
        pares = [f"{colname(pr.get('parentColumnId'))} → {colname(pr.get('childColumnId'))}"
                 for pr in (r.get("pairs") or [])]
        twin = any(rid2 in r_active and rels[rid2].get("parentTableId") == pt
                   and rels[rid2].get("childTableId") == ct for rid2 in rels if rid2 != rid)
        erw = r.get("erwinLongId")
        tag_of = rel_file.get(erw)
        xml_rel = models[tag_of].relationships.get(erw) if tag_of else None
        rows.append({
            "tipo": "Vínculo con columna no vigente", "_orden": 0, "_n": len(canvs),
            "relacion_erwin": xml_rel.name if xml_rel else "", "archivo": tag_of or "",
            "proyecto": proj, "tabla_padre": tname(pt), "padre_logico": tlogical(pt),
            "tabla_hija": tname(ct), "hija_logico": tlogical(ct),
            "pares": "\n".join(pares),
            "detalle": ("El vínculo usa una columna que no quedó vigente al unificar la tabla "
                        "entre archivos (ver hoja «Columnas distintas entre XML»)."
                        + (" Existe otra relación vigente entre las mismas tablas."
                           if twin else " No hay otra relación vigente entre las mismas tablas.")),
            "estado": "Desactivada: no se dibuja en el canvas",
            "canvases": cap_paths(canvs, "(las tablas no comparten canvas)"),
        })
    #  b) llave incompleta (ROJO)
    for rid in r_active:
        r = rels[rid]
        pt, ct = r.get("parentTableId"), r.get("childTableId")
        pk_ids = [cid for _pos, cid in pk_of_table.get(pt, [])]
        in_pairs = {pr.get("parentColumnId") for pr in (r.get("pairs") or [])}
        missing = [cid for cid in pk_ids if cid not in in_pairs]
        extra = [pr.get("parentColumnId") for pr in (r.get("pairs") or [])
                 if pr.get("parentColumnId") not in set(pk_ids)]
        if not missing and not extra:
            continue
        proj, canvs = rel_canvases(pt, ct)
        erw = r.get("erwinLongId")
        tag_of = rel_file.get(erw)
        xml_rel = models[tag_of].relationships.get(erw) if tag_of else None
        det = []
        if r.get("subcategory"):
            det.append("Relación de subcategoría.")
        if not pk_ids:
            det.append("La tabla padre NO tiene llave primaria definida.")
        if missing:
            det.append("PK del padre sin cubrir: "
                       + ", ".join(colname(cid, False) for cid in missing))
        if extra:
            det.append("Columnas conectadas que NO son PK del padre: "
                       + ", ".join(colname(cid, False) for cid in extra))
        rows.append({
            "tipo": "Llave incompleta (ROJO)", "_orden": 1, "_n": len(canvs),
            "relacion_erwin": xml_rel.name if xml_rel else "", "archivo": tag_of or "",
            "proyecto": proj, "tabla_padre": tname(pt), "padre_logico": tlogical(pt),
            "tabla_hija": tname(ct), "hija_logico": tlogical(ct),
            "pares": "\n".join(
                f"{colname(pr.get('parentColumnId'), False)} → {colname(pr.get('childColumnId'), False)}"
                for pr in (r.get("pairs") or [])),
            "detalle": "\n".join(det),
            "estado": "Vigente: la línea se pinta de ROJO en el canvas",
            "canvases": cap_paths(canvs, "(las tablas no comparten canvas)"),
        })
    #  c) no migradas (rotas de origen / sin pares de columnas)
    db_rel_erwin = {r.get("erwinLongId") for r in rels.values()}
    reused_by_file = {tag: {e["key"] for e in reports[tag]["decisions"]["rels_reused"]}
                      for tag in models}
    for tag, m in models.items():
        proj, dom = proj_of_tag[tag], dom_of_tag[tag]
        sof = schema_of[tag]
        for r in m.relationships.values():
            broken, motivo = False, ""
            if r.rel_type == ep.REL_TABLE_TO_VIEW:
                if r.parent_ref not in m.entities or r.child_ref not in m.views:
                    broken = True
                    partes = []
                    if r.parent_ref not in m.entities:
                        partes.append("el extremo padre no es una tabla del XML" +
                                      (" (es una vista)" if r.parent_ref in m.views else " (no existe)"))
                    if r.child_ref not in m.views:
                        partes.append("el extremo hijo no existe como vista en el XML")
                    motivo = "Relación tabla→vista rota: " + " y ".join(partes)
                else:
                    continue
            elif r.rel_type not in (ep.REL_IDENTIFYING, ep.REL_NON_IDENTIFYING, ep.REL_SUBTYPE):
                continue
            elif r.parent_ref not in m.entities or r.child_ref not in m.entities:
                broken = True
                motivo = "Relación con un extremo inexistente en el XML (rota de origen)"
            if not broken:
                if r.id in db_rel_erwin or r.name in reused_by_file[tag]:
                    continue
                motivo = ("Sin pares de columnas FK resolubles: no une columna con columna "
                          "(la plataforma exige columna origen y destino)")

            def _nm(ref):
                if ref in m.entities:
                    e = m.entities[ref]
                    return f"{pol.schema_or_default(sof.get(ref))}.{e.physical}", e.name
                if ref in m.views:
                    return f"{pol.schema_or_default(sof.get(ref))}.{m.views[ref].name} (vista)", ""
                return "(no existe en el XML)", ""

            pn, pl = _nm(r.parent_ref)
            cn, cl = _nm(r.child_ref)
            diags = rel_diags[tag].get(r.id) or []
            if not diags:
                diags = sorted(set(obj_diags[tag].get(r.parent_ref, []))
                               & set(obj_diags[tag].get(r.child_ref, [])))
            rows.append({
                "tipo": ("Rota de origen" if broken else "Sin pares de columnas"),
                "_orden": 2 if broken else 3, "_n": 0,
                "relacion_erwin": r.name, "archivo": tag, "proyecto": proj,
                "tabla_padre": pn, "padre_logico": pl, "tabla_hija": cn, "hija_logico": cl,
                "pares": "" if broken else "(no une columnas)",
                "detalle": motivo + f" — tipo {REL_TYPE_ES.get(r.rel_type, r.rel_type)}",
                "estado": "No existe en la plataforma",
                "canvases": cap_paths(
                    [f"{dom} / {sa} / {dg}" for sa, dg in sorted(set(diags))],
                    "(no está dibujada en ningún diagrama del XML)"),
            })
    rows.sort(key=lambda x: (x["_orden"], -x["_n"], x["tabla_padre"], x["tabla_hija"]))
    OUT["relaciones"] = rows
    log("[relaciones] " + "; ".join(
        f"{t}: {sum(1 for r in rows if r['tipo'] == t)}"
        for t in ["Vínculo con columna no vigente", "Llave incompleta (ROJO)",
                  "Rota de origen", "Sin pares de columnas"]))

    # ═══ Particiones con orden ilógico (recalculadas del XML) ══════════════
    grouped: dict[tuple, dict] = {}
    for tag, m in models.items():
        sof = schema_of[tag]
        uv = udp_vals_by_tag[tag]
        for e in m.entities.values():
            # columnas homónimas dentro de la tabla: vale el correlativo de la
            # copia que lo tenga (la migración conserva la copia con UDPs)
            by_name: dict[str, list] = {}
            for a in sorted(e.attributes, key=lambda a: a.order or 0):
                pn = (a.physical or "").upper()
                if pn:
                    by_name.setdefault(pn, []).append(a)
            vals = []
            for copies in by_name.values():
                for a in copies:
                    corr = pol.partition_correlative(uv.get((a.id, pol.PARTITION_UDP_KEY)))
                    if corr is not None:
                        vals.append((copies[0].order or 0, a, corr))
                        break
            vals = [(a, corr) for _o, a, corr in sorted(vals, key=lambda x: x[0])]
            if not vals:
                continue
            nums = [n for _a, n in vals]
            issues = []
            dups = sorted({n for n in nums if nums.count(n) > 1})
            if dups:
                issues.append("orden repetido: " + ", ".join(
                    f"PART_{n:02d} aparece {nums.count(n)} veces" for n in dups))
            missing = sorted(set(range(1, max(nums) + 1)) - set(nums))
            if missing:
                issues.append("orden salteado: falta " + ", ".join(
                    f"PART_{n:02d}" for n in missing))
            if nums != sorted(nums):
                issues.append("el correlativo no sigue el orden físico de las columnas")
            if not issues:
                continue
            nat_key = f"{pol.schema_or_default(sof.get(e.id))}.{e.physical}"
            problema = "; ".join(issues)
            detalle = "\n".join(f"{a.physical} → PART_{n:02d}" for a, n in vals[:15])
            if len(vals) > 15:
                detalle += f"\n… (+{len(vals) - 15} columnas más)"
            k = (nat_key.upper(), problema, detalle)
            if k in grouped:
                grouped[k]["tags"].add(tag)
                continue
            s, p = nat_key.upper().split(".", 1)
            tid = by_erwin_t.get(e.id) or by_key.get((s, p)) or by_phys.get(p)
            t = tables.get(tid, {})
            pj, canv = table_loc(tid, proj_of_tag[tag]) if tid else (proj_of_tag[tag], "?")
            grouped[k] = {"tags": {tag}, "proyecto": pj, "esquema": t.get("schema") or
                          nat_key.split(".", 1)[0], "tabla": e.physical,
                          "problema": problema, "detalle": detalle, "canvases": canv}
    rows = [{"archivos": ", ".join(sorted(g["tags"])), **{k: g[k] for k in
             ("proyecto", "esquema", "tabla", "problema", "detalle", "canvases")}}
            for g in grouped.values()]
    rows.sort(key=lambda x: (x["esquema"], x["tabla"]))
    OUT["particiones"] = rows
    log(f"[particiones] tablas con orden ilógico: {len(rows)}")

    # ═══ Vistas con observaciones ══════════════════════════════════════════
    rows = []
    for tag, m in models.items():          # a) descartadas
        proj, dom = proj_of_tag[tag], dom_of_tag[tag]
        sof = schema_of[tag]
        name_to_view = defaultdict(list)
        for v in m.views.values():
            name_to_view[v.name].append(v)
        for d in reports[tag]["decisions"]["views_discarded"]:
            diags, schema = [], ""
            for v in name_to_view.get(d["view"], []):
                schema = pol.schema_or_default(sof.get(v.id))
                diags += obj_diags[tag].get(v.id, [])
            rows.append({
                "archivo": tag, "proyecto": proj,
                "caso": "Sin tabla fuente resoluble (no migrada)",
                "esquema": schema, "vista": d["view"],
                "detalle": "No existe en la plataforma: ninguna de sus fuentes se pudo "
                           "resolver a una tabla física",
                "canvases": cap_paths([f"{dom} / {sa} / {dg}" for sa, dg in sorted(set(diags))],
                                      "(no está dibujada en ningún diagrama del XML)"),
            })
    n_dup_copias = 0
    for tag, m in models.items():          # b) duplicadas → queda una sola
        proj = proj_of_tag[tag]
        sof = schema_of[tag]
        groups = defaultdict(list)
        for v in m.views.values():
            groups[(pol.schema_or_default(sof.get(v.id)).upper(), v.name.upper())].append(v)
        for (sch, nm), copies in sorted(groups.items()):
            if len(copies) < 2:
                continue
            n_dup_copias += len(copies) - 1
            vdoc = next((v for v in views.values()
                         if (v.get("schema") or "").upper() == sch
                         and (v.get("name") or "").upper() == nm), None)
            canv = cap_paths(paths_of(layout_of_node.get(vdoc["_id"], []))) if vdoc else "?"
            rows.append({
                "archivo": tag, "proyecto": proj, "caso": "Duplicada en el archivo",
                "esquema": pol.schema_or_default(sof.get(copies[0].id)),
                "vista": copies[0].name,
                "detalle": f"Venía {len(copies)} veces en el archivo; en la plataforma queda "
                           f"una sola (la copia más usada/completa)",
                "canvases": canv,
            })
    rel_pairs_set = {frozenset((rels[rid].get("parentTableId"), rels[rid].get("childTableId")))
                     for rid in r_active}
    n_multi = 0
    for vid, v in views.items():           # c) multi-fuente sin join (audit C5)
        alive = [s for s in (v.get("sourceTableIds") or []) if s in t_active]
        if len(alive) <= 1 or (v.get("joinOverride") or "").strip():
            continue
        if all(any(frozenset((a, b)) in rel_pairs_set for b in alive if b != a) for a in alive):
            continue
        n_multi += 1
        fuentes = "\n".join(sorted(tname(s) for s in alive))
        tag_of = (view_file.get(v.get("erwinLongId")) or ["?"])[0]
        rows.append({
            "archivo": tag_of, "proyecto": project_of(layout_of_node.get(vid, [])),
            "caso": "Multi-fuente sin join declarado",
            "esquema": v.get("schema") or "", "vista": v.get("name") or "",
            "detalle": "Tiene 2+ tablas fuente sin join declarado ni relación entre ellas; "
                       "declarar el join en el editor de vistas\nFuentes:\n" + fuentes,
            "canvases": cap_paths(paths_of(layout_of_node.get(vid, []))),
        })
    n_vcol = 0
    for tag, m in models.items():          # d) columnas de vista sin origen
        proj, dom = proj_of_tag[tag], dom_of_tag[tag]
        sof = schema_of[tag]
        for v in m.views.values():
            for a in v.attributes:
                if a.parent_attr_ref:
                    continue
                n_vcol += 1
                rows.append({
                    "archivo": tag, "proyecto": proj, "caso": "Columna de vista sin origen",
                    "esquema": pol.schema_or_default(sof.get(v.id)),
                    "vista": f"{v.name}.{a.physical}",
                    "detalle": "La columna no pudo vincularse a su columna de origen; "
                               "se migró con definición propia, sin vínculo",
                    "canvases": cap_paths(
                        [f"{dom} / {sa} / {dg}"
                         for sa, dg in sorted(set(obj_diags[tag].get(v.id, [])))], "?"),
                })
    OUT["vistas"] = rows
    OUT["_vistas_nums"] = {"dup_copias": n_dup_copias, "multi": n_multi, "vcol": n_vcol}
    log(f"[vistas] con observaciones: {len(rows)}")

    # ═══ Diagramas homónimos con figuras distintas ═════════════════════════
    diag_versions: dict[str, dict[str, set]] = {}
    diag_meta: dict[str, tuple] = {}
    for tag, m in models.items():
        for d in m.diagrams:
            names = set()
            for ref, _a in d.shapes:
                if ref in m.entities:
                    names.add(m.entities[ref].physical or "?")
                elif ref in m.views:
                    names.add((m.views[ref].name or "?") + " (vista)")
            diag_versions.setdefault(d.id, {})[tag] = names
            diag_meta.setdefault(d.id, (d.subject_area, d.name))
    sa_by_erwin = {s.get("erwinLongId"): sid for sid, s in sas.items() if s.get("erwinLongId")}
    file_order = list(models)              # orden de carga de la corrida
    rows = []
    for did, per_tag in diag_versions.items():
        if len(per_tag) < 2 or len({frozenset(s) for s in per_tag.values()}) == 1:
            continue
        sid = sa_by_erwin.get(did)
        proj, path = sa_path.get(sid, ("?", "(canvas no encontrado)"))
        plat_tables = {tables[t].get("physicalName") for t in
                       ((sas[sid].get("tableIds") or []) if sid else []) if t in tables}
        det = []
        for tag in file_order:
            if tag not in per_tag:
                continue
            only = per_tag[tag] - set().union(*(s for t2, s in per_tag.items() if t2 != tag))
            det.append(f"{tag}: dibuja {len(per_tag[tag])} figuras"
                       + (f"; solo aquí: {cap_names(only)}" if only else ""))
        matched = [tag for tag in file_order if tag in per_tag and
                   {n for n in per_tag[tag] if not n.endswith(" (vista)")} == plat_tables]
        det.append(f"En la plataforma el canvas quedó con la versión del {matched[-1]}."
                   if matched else
                   f"En la plataforma el canvas tiene {len(plat_tables)} tablas.")
        rows.append({"tipo": "Mismo diagrama con figuras distintas", "proyecto": proj,
                     "objeto": f"«{diag_meta[did][0]} / {diag_meta[did][1]}»",
                     "archivos": ", ".join(t for t in file_order if t in per_tag),
                     "detalle": "\n".join(det), "canvas": path})
    #  tablas que quedaron sin figura en NINGÚN canvas por esa unificación
    n_diag_diff = len(rows)
    for tid in sorted(t_active, key=lambda x: tname(x)):
        if canvases_of_table.get(tid):
            continue
        t = tables[tid]
        erw = t.get("erwinLongId")
        tags = sorted(set(ent_file.get(erw, [])))
        drawn = [(g, sa, dg) for g in tags for sa, dg in obj_diags[g].get(erw, [])]
        if not drawn:
            continue                       # tampoco estaba dibujada en Erwin: no es incongruencia
        donde = "; ".join(f"«{sa} / {dg}» del {g}" for g, sa, dg in sorted(set(drawn)))
        rows.append({"tipo": "Tabla quedó sin figura", "proyecto": ", ".join(
                        sorted({proj_of_tag[g] for g in tags})),
                     "objeto": f"{t.get('schema')}.{t.get('physicalName')}",
                     "archivos": ", ".join(tags),
                     "detalle": (f"En Erwin está dibujada en {donde}, pero el canvas quedó con "
                                 f"la versión del otro archivo y la tabla no aparece en ningún "
                                 f"canvas de la plataforma. Si debe verse, volver a arrastrarla "
                                 f"al canvas (la tabla existe completa; se encuentra por el "
                                 f"buscador)."),
                     "canvas": "(ninguno)"})
    OUT["diagramas"] = rows
    log(f"[diagramas] con figuras distintas: {n_diag_diff}; "
        f"tablas sin figura por unificación: {len(rows) - n_diag_diff}")

    # ═══ Esquema no definido (archivos DDV) ════════════════════════════════
    rows = []
    n_t = n_v_desc = 0
    vby_erwin = {v.get("erwinLongId"): vid for vid, v in views.items()}
    vby_name = {}
    for vid, v in views.items():
        vby_name.setdefault((v.get("name") or "").upper(), vid)
    for tag, m in models.items():
        if not proj_of_tag[tag].upper().endswith("DDV"):
            continue
        sof = schema_of[tag]
        for e in sorted(m.entities.values(), key=lambda e: e.physical or ""):
            if e.id in sof:
                continue
            tid = by_erwin_t.get(e.id) or by_phys.get((e.physical or "").upper())
            if not tid or (tables[tid].get("schema") or "").upper() != pol.NO_SCHEMA.upper():
                continue
            t = tables[tid]
            p, canv = table_loc(tid, proj_of_tag[tag])
            rows.append({"tipo": "Tabla", "proyecto": p or proj_of_tag[tag],
                         "nombre": t.get("physicalName") or "",
                         "nombre_logico": t.get("logicalName") or "", "canvases": canv})
            n_t += 1
        for xv in sorted(m.views.values(), key=lambda v: v.name or ""):
            if xv.id in sof:
                continue
            vid = vby_erwin.get(xv.id) or vby_name.get((xv.name or "").upper())
            if not vid:
                n_v_desc += 1
                continue
            v = views[vid]
            if (v.get("schema") or "").upper() != pol.NO_SCHEMA.upper():
                continue
            rows.append({"tipo": "Vista",
                         "proyecto": project_of(layout_of_node.get(vid, []), proj_of_tag[tag]),
                         "nombre": v.get("name") or "", "nombre_logico": "",
                         "canvases": cap_paths(paths_of(layout_of_node.get(vid, [])))})
    OUT["no_definido"] = rows
    OUT["_no_definido_nums"] = {"tablas": n_t, "vistas": len(rows) - n_t,
                                "vistas_descartadas": n_v_desc}
    log(f"[esquema] no definido (DDV): {n_t} tablas + {len(rows) - n_t} vistas")

    # ═══ UDP incongruentes ═════════════════════════════════════════════════
    rows = []
    #  a) definiciones que no conversan entre archivos
    all_keys = sorted(set().union(*(collapsed_by_tag[tag].keys() for tag in models)))
    n_absent = n_diffdef = 0
    for key in all_keys:
        entries = {tag: collapsed_by_tag[tag][key] for tag in models
                   if key in collapsed_by_tag[tag]}
        present = [t for t in file_order if t in entries]
        absent = [t for t in file_order if t not in entries]
        any_e = entries[present[0]]
        nivel = NIVEL_ES.get(any_e["level"], any_e["level"]) + (
            " (lógico)" if any_e.get("view") == "logical" else "")   # doc 69: faceta
        resumen = any_e["dataType"] + (
            f", catálogo de {len(any_e['allowed_values'])} valores"
            if any_e["dataType"] == "list" and any_e["allowed_values"] else "")
        if absent:
            n_absent += 1
            rows.append({"tipo": "No está en todos los archivos", "udp": any_e["name"],
                         "nivel": nivel,
                         "archivos": f"Está en: {', '.join(present)}\nFalta en: {', '.join(absent)}",
                         "detalle": resumen, "objeto": "", "canvases": ""})
        diffs = []
        if len({e["dataType"] for e in entries.values()}) > 1:
            diffs.append("tipo de dato distinto: " + "; ".join(
                f"{t}: {entries[t]['dataType']}" for t in present))
        if any(e["dataType"] == "list" for e in entries.values()):
            avsets = {t: frozenset(v.strip().upper() for v in entries[t]["allowed_values"])
                      for t in present}
            if len(set(avsets.values())) > 1:
                diffs.append("catálogo de valores distinto:\n" + "\n".join(
                    f"  {t}: {cap_names(entries[t]['allowed_values']) or '(sin lista)'}"
                    for t in present))
        defaults = {t: (entries[t]["default"] or "").strip() for t in present}
        if len(set(defaults.values())) > 1:
            diffs.append("valor por defecto distinto: " + "; ".join(
                f"{t}: «{defaults[t] or '(vacío)'}»" for t in present))
        if diffs:
            n_diffdef += 1
            rows.append({"tipo": "Definición distinta entre archivos", "udp": any_e["name"],
                         "nivel": nivel, "archivos": ", ".join(present),
                         "detalle": "\n".join(diffs), "objeto": "", "canvases": ""})
    #  b) valores fuera de catálogo (en la plataforma, tal como venían)
    list_defs = {did: d for did, d in udp_defs.items()
                 if d.get("dataType") == "list" and d.get("allowedValues")}

    def _catalogo(d, val):
        return next((av for av in d["allowedValues"]
                     if av.strip().lower() == (val or "").strip().lower()), "")

    n_fuera = 0

    def _fuera(nivel, objeto, k, val, canv):
        nonlocal n_fuera
        d = list_defs.get(k)
        if not d or val in d["allowedValues"]:
            return
        n_fuera += 1
        cat = _catalogo(d, val)
        rows.append({"tipo": "Valor fuera de catálogo", "udp": d.get("name"), "nivel": nivel,
                     "archivos": "",
                     "detalle": f"valor guardado «{val}» vs catálogo «{cat or '(sin equivalente)'}»",
                     "objeto": objeto, "canvases": canv})

    for tid in t_active:
        for k, val in (tables[tid].get("udpValues") or {}).items():
            _fuera("Tabla", tname(tid), k, val, table_loc(tid)[1])
    for sid, s in sas.items():
        for k, val in (s.get("udpValues") or {}).items():
            _fuera("Canvas", f"(canvas) {sa_path[sid][1]}", k, val, sa_path[sid][1])
    for c in db.canonical_columns.find({**ACTIVE, "udpValues": {"$exists": True, "$ne": {}}},
                                       {"tableId": 1, "physicalName": 1, "udpValues": 1}):
        for k, val in (c.get("udpValues") or {}).items():
            _fuera("Columna", f"{tname(c.get('tableId'))}.{c.get('physicalName')}",
                   k, val, table_loc(c.get("tableId"))[1])
    #  c) definiciones sin ningún uso (no migradas)
    agg: dict = {}
    for tag in models:
        for d in reports[tag]["decisions"]["udp_defs_skipped"]:
            key = (d["name"], d["level"])
            agg.setdefault(key, {"name": d["name"], "level": d["level"],
                                 "dataType": d.get("dataType") or "", "files": []})
            agg[key]["files"].append(tag)
    for v in sorted(agg.values(), key=lambda x: x["name"]):
        rows.append({"tipo": "Definición sin ningún uso", "udp": v["name"],
                     "nivel": NIVEL_ES.get(v["level"], v["level"]),
                     "archivos": ", ".join(v["files"]),
                     "detalle": (v["dataType"] + " — " if v["dataType"] else "")
                     + "sin ningún valor en su archivo: no se migró",
                     "objeto": "", "canvases": ""})
    OUT["udp"] = rows
    log(f"[udp] no está en todos: {n_absent}; definición distinta: {n_diffdef}; "
        f"valor fuera de catálogo: {n_fuera}; sin uso: {len(agg)}")

    # ═══ dominios en conflicto (doc 75: unión distinta por proyecto) ══════
    rows = []
    for tag in models:
        for d in reports[tag]["decisions"].get("domain_conflicts", []):
            rows.append({"proyecto": proj_of_tag[tag], "dominio": d["name"],
                         "conservado": d["kept"], "ignorado": d["ignored"], "archivo": tag})
    rows.sort(key=lambda r: (r["proyecto"], r["dominio"].lower(), r["archivo"]))
    OUT["dominios"] = rows
    log(f"[dominios] en conflicto (tipo distinto entre archivos del mismo proyecto): {len(rows)}")

    # ═══ meta (solo datos de los archivos, sin narrativa) ══════════════════
    OUT["meta"] = {
        "archivos": [{
            "tag": fi["tag"], "xml": os.path.basename(fi["xml_path"]),
            "destino": f"{fi['project']} → {fi['domain']}",
            "tablas": fi["report"]["stats"].get("tablas", 0),
            "columnas": fi["report"]["stats"].get("columnas", 0),
            "vistas": fi["report"]["stats"].get("vistas", 0),
            "relaciones": fi["report"]["stats"].get("relaciones", 0),
            "canvases": fi["report"]["stats"].get("canvases", 0),
        } for fi in files],
    }
    return OUT


# ═══════════════════════════ Excel ═════════════════════════════════════════
def build_excel(D: dict, out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    AZUL = "1F4E78"
    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    HDR_FILL = PatternFill("solid", fgColor=AZUL)
    TITLE_FONT = Font(bold=True, size=14, color=AZUL)
    DESC_FONT = Font(italic=True, size=10, color="555555")
    THIN = Side(style="thin", color="C9D2DC")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP_TOP = Alignment(wrap_text=True, vertical="top")
    TOP = Alignment(vertical="top")
    CENTER = Alignment(horizontal="center", vertical="top")
    TAB = {"master": "C00000", "tabla": "2E75B6", "relacion": "C0504D",
           "vista": "31859C", "otro": "7F7F7F"}

    wb = Workbook()
    wb.remove(wb.active)

    def add_sheet(name, tab_kind, title, desc, headers, widths, rows, keys,
                  wrap_cols=(), num_cols=(), center_cols=(), max_h=150):
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = TAB[tab_kind]
        last = get_column_letter(len(headers))
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws.merge_cells(f"A2:{last}2")
        c = ws["A2"]
        c.value = desc
        c.font = DESC_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 14 * (desc.count("\n") + max(1, len(desc) // 160) + 1)
        hrow = 4
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=hrow, column=j, value=h)
            cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[hrow].height = 30
        for i, r in enumerate(rows):
            row_idx = hrow + 1 + i
            lines = 1
            for j, k in enumerate(keys, 1):
                v = r.get(k, "")
                cell = ws.cell(row=row_idx, column=j, value=v)
                cell.border = BORDER
                if j in num_cols or j in center_cols:
                    cell.alignment = CENTER
                elif j in wrap_cols:
                    cell.alignment = WRAP_TOP
                    if isinstance(v, str):
                        lines = max(lines, v.count("\n") + 1)
                else:
                    cell.alignment = TOP
            if lines > 1:
                ws.row_dimensions[row_idx].height = min(13.5 * lines + 4, max_h)
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        if rows:
            ws.auto_filter.ref = f"A{hrow}:{last}{hrow + len(rows)}"
            ws.conditional_formatting.add(
                f"A{hrow + 1}:{last}{hrow + len(rows)}",
                FormulaRule(formula=["MOD(ROW(),2)=0"],
                            fill=PatternFill("solid", fgColor="F2F5F9")))
        ws.freeze_panes = f"A{hrow + 1}"

    # ── hoja maestra «Incongruencias» ──────────────────────────────────────
    vn = D["_vistas_nums"]
    nd = D["_no_definido_nums"]
    rel_counts = defaultdict(int)
    for r in D["relaciones"]:
        rel_counts[r["tipo"]] += 1
    udp_counts = defaultdict(int)
    for r in D["udp"]:
        udp_counts[r["tipo"]] += 1
    n_views_disc = sum(1 for r in D["vistas"] if r["caso"].startswith("Sin tabla fuente"))
    n_views_dup = sum(1 for r in D["vistas"] if r["caso"].startswith("Duplicada"))

    MASTER = [
        ("Tablas homónimas", len(D["homonimas"]), "Tablas homónimas (_DUPn)",
         "El mismo nombre físico de tabla se usa para tablas DISTINTAS (columnas diferentes) "
         "en uno o varios modelos.",
         "Definir un nombre único por tabla, o igualar las estructuras si en realidad son la "
         "misma tabla.",
         "La copia más usada conservó el nombre; las demás llevan sufijo _DUPn con todo su "
         "contenido intacto."),
        ("Misma tabla con columnas distintas", len(D["cols_distintas"]),
         "Columnas distintas entre XML",
         "La misma tabla no tiene las mismas columnas en todos los archivos donde aparece "
         "(cada fila es una columna que está en unos archivos y falta en otros).",
         "Hacer merge de columnas en Erwin: igualar la estructura de la tabla en todos los "
         "modelos donde aparece.",
         "Quedó la versión más usada de la tabla; las columnas de la otra versión no están "
         "vigentes en la plataforma."),
        ("Columnas duplicadas dentro de una tabla", len(D["cols_duplicadas"]),
         "Columnas duplicadas",
         "La tabla trae la misma columna dos o más veces, idéntica, en el propio archivo.",
         "Eliminar la columna repetida en Erwin.",
         "Se conservó una sola copia por tabla (las duplicadas son lo único que se eliminó)."),
        ("Relaciones incongruentes",
         len(D["relaciones"]), "Relaciones incongruentes",
         f"Vínculo con columna no vigente: {rel_counts['Vínculo con columna no vigente']} · "
         f"llave incompleta (ROJO): {rel_counts['Llave incompleta (ROJO)']} · "
         f"rotas de origen: {rel_counts['Rota de origen']} · "
         f"sin pares de columnas: {rel_counts['Sin pares de columnas']}.",
         "Revisar la llave primaria del padre y las columnas del vínculo; volver a trazar las "
         "rotas; definir columna origen y destino en las que no unen columnas.",
         "Las de llave incompleta se pintan de ROJO en el canvas; las de columna no vigente "
         "quedaron desactivadas; las rotas o sin pares no se migraron."),
        ("Particiones con orden ilógico", len(D["particiones"]),
         "Particiones con orden ilógico",
         "Los correlativos de partición (PART_01, PART_02, …) vienen repetidos, salteados o "
         "en desorden respecto al orden físico de las columnas.",
         "Corregir la numeración PART_nn en la UDP «Particion» de cada columna.",
         "El particionado quedó por el ORDEN FÍSICO de las columnas (el correlativo original "
         "no se tocó)."),
        ("Vistas con observaciones", len(D["vistas"]), "Vistas",
         f"Sin tabla fuente resoluble: {n_views_disc} · duplicadas en el archivo: "
         f"{n_views_dup} ({vn['dup_copias']} copias) · multi-fuente sin join: {vn['multi']} · "
         f"columnas de vista sin origen: {vn['vcol']}.",
         "Revisar las fuentes de cada vista y el vínculo de sus columnas.",
         "Las sin fuente no se migraron; de las duplicadas quedó una; el join pendiente se "
         "declara en el editor de vistas."),
        ("Objetos sin esquema (DDV)", len(D["no_definido"]), "Esquema no definido",
         f"Tablas ({nd['tablas']}) y vistas ({nd['vistas']}) de los archivos DDV sin "
         f"Hive_Database (esquema) en el XML.",
         "Asignar el Hive_Database real a cada objeto en Erwin.",
         "Entraron al esquema «No_Definido» (en el UDV no manejar esquema es lo esperado y "
         "no se lista)."),
        ("UDP que no conversan entre archivos", len(D["udp"]), "UDP incongruentes",
         f"No está en todos los archivos: {udp_counts['No está en todos los archivos']} · "
         f"definición distinta: {udp_counts['Definición distinta entre archivos']} · "
         f"valores fuera de catálogo: {udp_counts['Valor fuera de catálogo']} · "
         f"definiciones sin uso: {udp_counts['Definición sin ningún uso']}.",
         "Homologar las definiciones UDP (nombre, tipo, catálogo de valores) entre los "
         "modelos, y corregir los valores que no calzan con el catálogo.",
         "Se migró la unión de definiciones; los valores se guardaron tal como venían."),
        ("Dominios en conflicto", len(D["dominios"]), "Dominios en conflicto",
         "El mismo parent domain viene con tipo físico DISTINTO en dos archivos del mismo "
         "proyecto (cada proyecto tiene sus propios dominios; entre sus archivos gana el primero).",
         "Igualar el tipo del dominio en Erwin en todos los modelos del proyecto.",
         "Se conservó el tipo del primer archivo cargado; el otro tipo quedó registrado y no se "
         "aplicó."),
        ("Diagramas homónimos con figuras distintas", len(D["diagramas"]),
         "Diagramas con figuras distintas",
         "El mismo diagrama (mismo identificador Erwin) existe en dos archivos con figuras "
         "diferentes; incluye las tablas que quedaron sin figura en ningún canvas.",
         "Igualar los diagramas homónimos entre modelos (o confirmar cuál versión es la "
         "válida).",
         "El canvas quedó con las figuras del último archivo cargado; las tablas afectadas "
         "existen completas y pueden volver a arrastrarse al canvas."),
    ]

    ws = wb.create_sheet("Incongruencias")
    ws.sheet_properties.tabColor = TAB["master"]
    ws.column_dimensions["A"].width = 3
    for col, w in (("B", 34), ("C", 9), ("D", 56), ("E", 46), ("F", 46), ("G", 30)):
        ws.column_dimensions[col].width = w
    r = 2
    ws.cell(row=r, column=2, value="Incongruencias encontradas en los archivos Erwin").font = \
        Font(bold=True, size=15, color=AZUL)
    r += 2
    for j, h in enumerate(["Incongruencia", "Casos", "Qué se encontró en los archivos XML",
                           "Limpieza sugerida en Erwin", "Qué hizo la migración en la plataforma",
                           "Detalle caso por caso"], 2):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 26
    r += 1
    for nombre, n, hoja, hallazgo, erwin, script in MASTER:
        vals = [(2, nombre, None), (3, n, "num"), (4, hallazgo, "wrap"),
                (5, erwin, "wrap"), (6, script, "wrap"), (7, hoja, "link")]
        h = max(len(hallazgo), len(erwin), len(script)) // 45 + 2
        for j, v, kind in vals:
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            if kind == "num":
                c.number_format = "#,##0"
                c.alignment = CENTER
                c.font = Font(bold=True, size=10)
            elif kind == "wrap":
                c.alignment = WRAP_TOP
                c.font = Font(size=9)
            elif kind == "link":
                c.hyperlink = f"#'{hoja}'!A1"
                c.font = Font(color="0563C1", underline="single", size=10)
                c.alignment = TOP
            else:
                c.font = Font(bold=True, size=10)
                c.alignment = WRAP_TOP
        ws.row_dimensions[r].height = 13 * h
        r += 1
    r += 1

    ws.cell(row=r, column=2, value="Archivos analizados").font = Font(bold=True, size=12, color=AZUL)
    r += 1
    hdr = ["Nombre corto", "Archivo XML", "Proyecto → carpeta", "Tablas", "Columnas",
           "Vistas", "Relaciones", "Canvases"]
    for j, h in enumerate(hdr, 2):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
    r += 1
    for fi in D["meta"]["archivos"]:
        vals = [fi["tag"], fi["xml"], fi["destino"], fi["tablas"], fi["columnas"],
                fi["vistas"], fi["relaciones"], fi["canvases"]]
        for j, v in enumerate(vals, 2):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            if isinstance(v, int):
                c.number_format = "#,##0"
        r += 1
    ws.sheet_view.showGridLines = False

    # ── hojas de detalle ───────────────────────────────────────────────────
    add_sheet(
        "Tablas homónimas (_DUPn)", "tabla",
        "Tablas homónimas: el mismo nombre para tablas distintas",
        "El nombre físico de tabla debe ser único en toda la plataforma. La copia MÁS USADA (más "
        "relaciones, vistas y diagramas que la referencian) conservó el nombre y las demás copias se "
        "renombraron con sufijo _DUPn, con todo su contenido intacto. Cada fila es una copia renombrada. "
        "Limpieza en Erwin: darle un nombre propio a cada tabla, o igualar las estructuras si son la misma.",
        ["Archivo XML", "Proyecto", "Esquema", "Nombre original", "Nombre lógico",
         "Nombre final en la plataforma", "Copias con ese nombre", "Copia que conservó el nombre",
         "Canvases donde aparece esta copia (_DUPn)", "Canvases de la copia que conservó el nombre"],
        [12, 12, 26, 32, 32, 36, 9, 32, 52, 52],
        D["homonimas"],
        ["archivo", "proyecto", "esquema", "nombre_original", "nombre_logico",
         "nombre_final", "copias", "conservo_nombre", "canvases", "canvases_kept"],
        wrap_cols=(9, 10), num_cols=(7,))

    add_sheet(
        "Columnas distintas entre XML", "tabla",
        "Misma tabla con columnas distintas según el archivo (merge pendiente en Erwin)",
        "La misma tabla aparece en dos o más archivos con columnas que no coinciden. Cada fila es una "
        "columna que está en unos archivos y falta en otros. En la plataforma quedó la versión MÁS USADA "
        "de la tabla: las columnas que solo estaban en la otra versión figuran como «No vigente». "
        "Limpieza en Erwin: hacer merge de columnas para que la tabla tenga la misma estructura en todos "
        "los modelos.",
        ["Proyecto", "Esquema", "Tabla", "Columna", "Nombre lógico", "¿Es PK?",
         "Está en", "Falta en", "En la plataforma", "Canvases de la tabla"],
        [12, 22, 34, 30, 30, 8, 22, 22, 30, 52],
        D["cols_distintas"],
        ["proyecto", "esquema", "tabla", "columna", "columna_logica", "es_pk",
         "presente", "falta", "estado", "canvases"],
        wrap_cols=(9, 10), center_cols=(6,))

    add_sheet(
        "Columnas duplicadas", "tabla",
        "Columnas duplicadas dentro de una misma tabla (eliminadas)",
        "La tabla traía la misma columna dos o más veces, idéntica (mismo nombre, tipo y definición). Se "
        "conservó una sola copia por tabla, priorizando la que participa en relaciones, luego la PK, la "
        "que tiene definición, dominio y valores UDP — este es el único caso de columnas eliminadas. "
        "Limpieza en Erwin: quitar la columna repetida.",
        ["Archivo XML", "Proyecto", "Esquema", "Tabla", "Columna duplicada", "Canvases de la tabla"],
        [12, 12, 26, 38, 30, 56],
        D["cols_duplicadas"],
        ["archivo", "proyecto", "esquema", "tabla", "columna", "canvases"],
        wrap_cols=(6,))

    add_sheet(
        "Relaciones incongruentes", "relacion",
        "Relaciones incongruentes (todas en una sola hoja — filtrar por «Tipo»)",
        "Cuatro tipos. «Vínculo con columna no vigente»: usaba una columna que no quedó vigente al "
        "unificar la tabla — está desactivada. «Llave incompleta (ROJO)»: las columnas conectadas no "
        "cubren la llave primaria completa del padre — la línea se pinta de ROJO en el canvas. «Rota de "
        "origen»: apunta a un objeto que no existe en el propio XML. «Sin pares de columnas»: no une "
        "columna con columna. Limpieza en Erwin: revisar llaves y columnas del vínculo en cada caso.",
        ["Tipo", "Relación en Erwin", "Archivo XML", "Proyecto", "Tabla padre", "Padre (lógico)",
         "Tabla hija", "Hija (lógico)", "Columnas del vínculo (padre → hija)", "Detalle",
         "Estado en la plataforma", "Canvases / dónde verla"],
        [24, 13, 12, 12, 32, 24, 32, 24, 42, 44, 24, 52],
        D["relaciones"],
        ["tipo", "relacion_erwin", "archivo", "proyecto", "tabla_padre", "padre_logico",
         "tabla_hija", "hija_logico", "pares", "detalle", "estado", "canvases"],
        wrap_cols=(9, 10, 11, 12))

    add_sheet(
        "Particiones con orden ilógico", "tabla",
        "Tablas cuyo correlativo de partición no es lógico",
        "Los valores PART_nn de la UDP «Particion» vienen con orden repetido, salteado o distinto al "
        "orden físico de las columnas. La columna «Columnas → correlativo» muestra lo leído del XML, en "
        "orden físico. En la plataforma el particionado quedó por el orden físico de las columnas. "
        "Limpieza en Erwin: renumerar PART_01, PART_02, … sin repetir ni saltear.",
        ["Archivos", "Proyecto", "Esquema", "Tabla", "Problema",
         "Columnas → correlativo (orden físico)", "Canvases de la tabla"],
        [16, 12, 24, 36, 44, 36, 52],
        D["particiones"],
        ["archivos", "proyecto", "esquema", "tabla", "problema", "detalle", "canvases"],
        wrap_cols=(5, 6, 7))

    add_sheet(
        "Vistas", "vista",
        "Vistas con observaciones (filtrar por «Caso»)",
        "Sin tabla fuente resoluble: ninguna fuente se pudo resolver a una tabla física — no existen en "
        "la plataforma. Duplicada en el archivo: mismo esquema y nombre repetido — quedó una sola. "
        "Multi-fuente sin join: 2+ tablas fuente sin join declarado ni relación entre ellas — declararlo "
        "en el editor de vistas. Columna de vista sin origen: no pudo vincularse a su columna fuente. "
        "Limpieza en Erwin: revisar fuentes y vínculos de estas vistas.",
        ["Caso", "Archivo XML", "Proyecto", "Esquema", "Vista", "Detalle", "Canvases / dónde verla"],
        [26, 12, 12, 28, 44, 52, 52],
        D["vistas"],
        ["caso", "archivo", "proyecto", "esquema", "vista", "detalle", "canvases"],
        wrap_cols=(6, 7))

    add_sheet(
        "Esquema no definido", "otro",
        "Tablas y vistas de los archivos DDV sin esquema (entraron a «No_Definido»)",
        "Estos objetos no traían Hive_Database (esquema) en el XML. A diferencia del UDV —donde no "
        "manejar esquema es lo esperado y por eso no se lista—, en el DDV probablemente sí deberían "
        "tener uno. Limpieza en Erwin: asignar el Hive_Database real a cada objeto (en la plataforma "
        "puede corregirse desde Properties).",
        ["Tipo", "Proyecto", "Nombre", "Nombre lógico", "Canvases donde aparece"],
        [8, 12, 44, 36, 60],
        D["no_definido"],
        ["tipo", "proyecto", "nombre", "nombre_logico", "canvases"],
        wrap_cols=(5,), center_cols=(1,))

    add_sheet(
        "UDP incongruentes", "otro",
        "Propiedades UDP que no conversan entre los archivos (filtrar por «Tipo»)",
        "«No está en todos los archivos»: la definición existe en unos XML y falta en otros. «Definición "
        "distinta entre archivos»: mismo nombre pero tipo de dato, catálogo de valores o default "
        "diferentes. «Valor fuera de catálogo»: el valor guardado no calza con el catálogo de su "
        "definición (difiere en mayúsculas/minúsculas). «Definición sin ningún uso»: no se migró. "
        "Limpieza en Erwin: homologar las definiciones UDP entre modelos y corregir los valores.",
        ["Tipo", "UDP", "Nivel", "Archivos", "Detalle", "Objeto afectado", "Canvases"],
        [26, 24, 10, 26, 48, 44, 44],
        D["udp"],
        ["tipo", "udp", "nivel", "archivos", "detalle", "objeto", "canvases"],
        wrap_cols=(4, 5, 6, 7), center_cols=(3,))

    add_sheet(
        "Diagramas con figuras distintas", "otro",
        "El mismo diagrama existe en dos archivos con figuras diferentes",
        "El diagrama comparte identificador Erwin entre archivos pero no dibuja las mismas figuras en "
        "cada uno; el canvas de la plataforma quedó con la versión del último archivo cargado. Incluye "
        "las tablas que por esa razón no aparecen en ningún canvas (existen completas: basta volver a "
        "arrastrarlas). Limpieza en Erwin: igualar los diagramas homónimos o confirmar cuál versión vale.",
        ["Tipo", "Proyecto", "Diagrama / tabla", "Archivos", "Detalle", "Canvas en la plataforma"],
        [28, 14, 40, 20, 70, 44],
        D["diagramas"],
        ["tipo", "proyecto", "objeto", "archivos", "detalle", "canvas"],
        wrap_cols=(5, 6))

    add_sheet(
        "Dominios en conflicto", "otro",
        "El mismo parent domain con tipo distinto entre archivos del proyecto",
        "Doc 75: los dominios son por proyecto y se forman como la unión DISTINTA de sus archivos: "
        "el primer archivo cargado define el tipo; un archivo posterior con OTRO tipo para el mismo "
        "dominio se registra acá y no se aplica. Limpieza en Erwin: igualar el tipo del dominio en "
        "todos los modelos del proyecto.",
        ["Proyecto", "Dominio", "Tipo conservado", "Tipo ignorado", "Archivo"],
        [18, 34, 26, 26, 28],
        D["dominios"],
        ["proyecto", "dominio", "conservado", "ignorado", "archivo"],
        center_cols=(3, 4))

    wb.save(out_path)
    log(f"OK → {out_path} ({os.path.getsize(out_path) // 1024} KB, {len(wb.sheetnames)} hojas)")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--reports-dir", default=os.path.join(BACKEND, "migration-reports"))
    ap.add_argument("--xml-root", default=os.path.join(BACKEND, "..", "folder_data"))
    ap.add_argument("--out", default=os.path.join(BACKEND, "..", "REPORTE-MIGRACION-ERWIN-DETALLE.xlsx"))
    ap.add_argument("--cache-dir", default=None,
                    help="Carpeta para cachear los XML parseados (pickle) entre corridas")
    args = ap.parse_args(argv)

    files = discover(args.reports_dir, args.xml_root)
    log("Archivos de la última corrida:")
    for fi in files:
        log(f"  · {fi['tag']}: {fi['xml_path']}")
    parse_models(files, args.cache_dir)
    data = extract(files)
    build_excel(data, os.path.abspath(args.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
