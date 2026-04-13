"""Script para generar el archivo sample_input.xlsx de ejemplo.

Crea un archivo Excel con 2 tablas relacionadas (customers y orders)
para testing del sistema Data Modeler Agent.
"""

import sys
from pathlib import Path

# Agregar raíz al path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "data" / "sample_input.xlsx"


def create_sample():
    """Genera el archivo sample_input.xlsx con tablas de ejemplo."""
    wb = Workbook()

    # ─── Pestaña 1: customers ──────────────────────────────────
    ws1 = wb.active
    ws1.title = "customers"

    headers = ["column_name", "functional_definition", "data_type_hint", "is_nullable", "notes"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    customers_data = [
        ["", "Identificador único del cliente", "BIGINT", "No", "Clave primaria"],
        ["nombre_cliente", "Nombre completo del cliente", "VARCHAR", "No", ""],
        ["", "Dirección de correo electrónico del cliente", "VARCHAR", "No", "Único"],
        ["telefono", "Número de teléfono de contacto del cliente", "VARCHAR", "Si", ""],
        ["", "Dirección física del cliente", "TEXT", "Si", "Dirección completa"],
        ["", "Fecha de nacimiento del cliente", "DATE", "Si", ""],
        ["", "Indicador booleano de si el registro está activo", "BOOLEAN", "No", "Default true"],
        ["segmento", "Segmento o categoría comercial del cliente", "VARCHAR", "Si", "Gold, Silver, Bronze"],
    ]

    for row_idx, row_data in enumerate(customers_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws1.cell(row=row_idx, column=col_idx, value=value)

    # Ajustar anchos
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 50
    ws1.column_dimensions["C"].width = 15
    ws1.column_dimensions["D"].width = 12
    ws1.column_dimensions["E"].width = 25

    # ─── Pestaña 2: orders ─────────────────────────────────────
    ws2 = wb.create_sheet("orders")

    for col_idx, header in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    orders_data = [
        ["", "Identificador único de la orden", "BIGINT", "No", "Clave primaria"],
        ["", "Identificador único del cliente", "BIGINT", "No", "FK hacia customers"],
        ["", "Fecha en que se realizó la orden", "TIMESTAMP", "No", ""],
        ["", "Código de estado del registro o transacción", "VARCHAR", "No", "pending, completed, cancelled"],
        ["", "Monto total de una transacción u orden", "DECIMAL", "No", ""],
        ["desc_orden", "Descripción o detalle de la orden", "TEXT", "Si", ""],
        ["", "Método de pago utilizado", "VARCHAR", "Si", "credit_card, debit, transfer"],
        ["", "Dirección de envío de la orden", "TEXT", "Si", ""],
        ["", "Fecha de ejecución de un proceso o rutina", "DATE", "Si", "Fecha de procesamiento batch"],
    ]

    for row_idx, row_data in enumerate(orders_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws2.cell(row=row_idx, column=col_idx, value=value)

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 50
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 12
    ws2.column_dimensions["E"].width = 35

    # ─── Guardar ───────────────────────────────────────────────
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"✓ Archivo creado: {OUTPUT_PATH}")
    print(f"  • Pestaña 'customers': {len(customers_data)} columnas")
    print(f"  • Pestaña 'orders': {len(orders_data)} columnas")


if __name__ == "__main__":
    create_sample()
