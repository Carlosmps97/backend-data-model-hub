"""Doc 105 (ronda 5, revisor R13) — el reporte de incongruencias y una relación
que el kit omitió porque una de sus tablas se BORRÓ en la plataforma.

El kit sólo sumaba la estadística «relaciones omitidas (tabla borrada en la
app)», sin decisión en el reporte; `migration_detail_report` la clasificaba
como «Sin pares de columnas FK resolubles…» y le pedía al modelador corregir en
Erwin una relación que está bien. Ahora el kit la registra
(`rels_deleted_in_app`) y el reporte la etiqueta «Tabla borrada en la
plataforma», con su leyenda y su conteo en el resumen — como las vistas."""
from __future__ import annotations

from openpyxl import load_workbook

import scripts.migration_detail_report as mdr
from scripts.erwin_migration.migrate import Migrator
from tests.erwin_migration.test_subcategory_projection_doc100 import ProjectingDb
from tests.erwin_migration.test_tablas_borradas_doc105 import _borrar_tabla_en_la_app, _modelo

GONE = "Tabla borrada en la plataforma"


def _corrida_con_tabla_borrada(monkeypatch) -> dict:
    db = ProjectingDb()
    primera = Migrator(db, _modelo(False), "Fam", None)
    primera.run()
    _borrar_tabla_en_la_app(db, primera.table_pid["E1"])
    model = _modelo(True)                                          # trae R19 TAB_UNO→TAB_NUEVE (nueva)
    mig = Migrator(db, model, "Fam", None)
    mig.run()
    assert mig.report["rels_deleted_in_app"] == [{"erwinLongId": "R19", "name": "R19"}]
    monkeypatch.setattr(mdr, "get_sync_db", lambda: db)
    monkeypatch.setattr(mdr, "log", lambda _m: None)
    return mdr.extract([{"tag": "A", "project": "Fam", "domain": "D", "model": model, "xml_path": "a.xml",
                         "report": {"stats": dict(mig.stats), "decisions": mig.report}}])


def test_la_relacion_omitida_por_tabla_borrada_se_etiqueta_como_tal(monkeypatch):
    out = _corrida_con_tabla_borrada(monkeypatch)
    (row,) = [r for r in out["relaciones"] if r["relacion_erwin"] == "R19"]
    assert row["tipo"] == GONE
    assert "borrada" in row["detalle"] and "Sin pares" not in row["detalle"]


def test_el_resumen_y_la_leyenda_cuentan_el_caso(monkeypatch, tmp_path):
    out = _corrida_con_tabla_borrada(monkeypatch)
    path = tmp_path / "reporte.xlsx"
    mdr.build_excel(out, str(path))
    wb = load_workbook(path)
    textos = [str(c.value) for ws in wb.worksheets for fila in ws.iter_rows() for c in fila if c.value]
    assert any("tabla borrada en la plataforma: 1" in t for t in textos)                   # resumen
    assert any(f"«{GONE}»" in t for t in textos)                                           # leyenda


def test_un_reporte_anterior_sin_la_clave_nueva_sigue_leyendose(monkeypatch):
    """Los JSON de corridas anteriores no traen `rels_deleted_in_app`."""
    db = ProjectingDb()
    model = _modelo(True)
    mig = Migrator(db, model, "Fam", None)
    mig.run()
    decisions = {k: v for k, v in mig.report.items() if k != "rels_deleted_in_app"}
    monkeypatch.setattr(mdr, "get_sync_db", lambda: db)
    monkeypatch.setattr(mdr, "log", lambda _m: None)
    out = mdr.extract([{"tag": "A", "project": "Fam", "domain": "D", "model": model, "xml_path": "a.xml",
                        "report": {"stats": dict(mig.stats), "decisions": decisions}}])
    assert all(r["tipo"] != GONE for r in out["relaciones"])
